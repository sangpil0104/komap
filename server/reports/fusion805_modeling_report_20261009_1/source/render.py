"""Render one structured manuscript to matching Markdown and native Word content."""
from pathlib import Path
import re,json,csv,sys
from datetime import datetime,timezone
from PIL import Image
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START,WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

FONT='NanumGothic'

def font(run,size=None,bold=None):
    run.font.name=FONT
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
    if size:run.font.size=Pt(size)
    if bold is not None:run.bold=bold
    return run

def hyperlink(p,text,url):
    item=OxmlElement('w:hyperlink');item.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');props=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'2455A4');props.append(c)
    fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),FONT);fonts.set(qn('w:eastAsia'),FONT);props.append(fonts)
    run.append(props);t=OxmlElement('w:t');t.text=text;run.append(t);item.append(run);p._p.append(item)

def inline(p,text,refs):
    for part in re.split(r'(\*\*.*?\*\*|`[^`]+`|\[S\d+\]|\[P\d+\])',str(text)):
        if not part:continue
        if re.fullmatch(r'\[(?:S|P)\d+\]',part) and part[1:-1] in refs:
            hyperlink(p,part,refs[part[1:-1]]['url'])
        elif part.startswith('**') and part.endswith('**'):font(p.add_run(part[2:-2]),bold=True)
        elif part.startswith('`') and part.endswith('`'):font(p.add_run(part[1:-1]),9)
        else:font(p.add_run(part))

def style_table(table):
    table.style='Table Grid';table.autofit=True
    for i,row in enumerate(table.rows):
        pr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');pr.append(cant)
        if i==0:
            repeat=OxmlElement('w:tblHeader');pr.append(repeat)
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                if len(table.rows)<=8:p.paragraph_format.keep_with_next=True
                for r in p.runs:font(r,8.5,i==0)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EFF8');cell._tc.get_or_add_tcPr().append(sh)

def render(package,blocks,refs):
    doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    for side in ['top_margin','bottom_margin','left_margin','right_margin']:setattr(sec,side,Cm(1.8))
    sec.header_distance=Cm(.8);sec.footer_distance=Cm(.8)
    normal=doc.styles['Normal'];normal.font.name=FONT;normal.font.size=Pt(10)
    normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
    normal.paragraph_format.line_spacing=1.22;normal.paragraph_format.space_after=Pt(7)
    for name,size in [('Title',24),('Heading 1',17),('Heading 2',13),('Heading 3',11)]:
        style=doc.styles[name];style.font.name=FONT;style.font.size=Pt(size);style.font.color.rgb=RGBColor.from_string('20324A')
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT);style.paragraph_format.space_before=Pt(15);style.paragraph_format.space_after=Pt(8)
    hp=sec.header.paragraphs[0];font(hp.add_run('KoMaP | 고탄성 알루미늄 합금 상 분할 · Fusion805'),8);hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER;font(fp.add_run('모델링 결과서  •  '),8)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld)
    doc.core_properties.title='고탄성 알루미늄 합금 이미지 상 분할 모델링 결과서 — Fusion805'
    doc.core_properties.subject='프로젝트 실측 결과와 GitHub 고정 모델을 구분한 확장 보고서'
    doc.core_properties.author='KoMaP 프로젝트'
    doc.core_properties.created=datetime.now(timezone.utc)
    doc.core_properties.modified=datetime.now(timezone.utc)
    md=[];table_catalog=[];fig_counter=0;table_counter=0
    for b in blocks:
        typ=b['type']
        if typ=='title':
            doc.add_paragraph(b['text'],'Title');md+=['# '+b['text'],'']
        elif typ=='subtitle':
            p=doc.add_paragraph();font(p.add_run(b['text']),14);md+=[b['text'],'']
        elif typ=='heading':
            if b.get('newpage'):doc.add_page_break()
            doc.add_heading(b['text'],b['level']);md+=['#'*(b['level']+1)+' '+b['text'],'']
        elif typ in ['p','equation']:
            p=doc.add_paragraph();inline(p,b['text'],refs)
            if typ=='equation':p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            md+=[b.get('markdown',b['text']),'']
        elif typ=='pagebreak':doc.add_page_break()
        elif typ=='table':
            table_counter+=1;caption=f'표 {table_counter}. {b["title"]}'
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;font(p.add_run(caption),10,True)
            table=doc.add_table(rows=1,cols=len(b['headers']))
            for c,v in zip(table.rows[0].cells,b['headers']):c.text=str(v)
            for row in b['rows']:
                for c,v in zip(table.add_row().cells,row):c.text=str(v)
            style_table(table)
            note=b.get('note','');p=doc.add_paragraph();inline(p,note,refs)
            for run in p.runs:font(run,8)
            lines=[f'**{caption}**','', '| '+' | '.join(map(str,b['headers']))+' |','| '+' | '.join(['---']*len(b['headers']))+' |']
            lines+=['| '+' | '.join(str(x).replace('|','/').replace('\n','<br>') for x in row)+' |' for row in b['rows']]
            lines+=['',note,''];md+=lines
            stem=f'{table_counter:02d}_{b["id"]}';(package/'tables'/f'{stem}.md').write_text('\n'.join(lines))
            with (package/'tables'/f'{stem}.csv').open('w',encoding='utf-8-sig',newline='') as f:
                w=csv.writer(f);w.writerow(b['headers']);w.writerows(b['rows'])
            table_catalog.append(dict(number=table_counter,id=b['id'],title=b['title'],file=f'tables/{stem}.md',rows=len(b['rows'])))
        elif typ=='figure':
            fig_counter+=1;path=package/b['file'];w,h=Image.open(path).size
            width=min(16.8,18*w/h)
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(path),width=Cm(width))
            caption=f'그림 {fig_counter}. {b["title"]}. {b["caption"]}'
            p=doc.add_paragraph();inline(p,caption,refs)
            for run in p.runs:font(run,8.5)
            md += [f'![그림 {fig_counter}. {b["title"]}]({b["file"]})','',caption,'']
        elif typ=='reference':
            p=doc.add_paragraph();font(p.add_run('['+b['id']+'] '),9,True);hyperlink(p,b['title'],b['url']);inline(p,' — '+b.get('note',''),refs)
            md += [f'[{b["id"]}] [{b["title"]}]({b["url"]}) — {b.get("note","")}','']
        elif typ=='appendix_start':
            sec=doc.add_section(WD_SECTION_START.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width=Cm(29.7);sec.page_height=Cm(21)
            for side in ['top_margin','bottom_margin','left_margin','right_margin']:setattr(sec,side,Cm(1.3))
            md+=['## 부록 A. 검증 20장 전체 비교','']
        elif typ=='appendix_image':
            if b['index']>1:doc.add_page_break()
            p=doc.add_paragraph();font(p.add_run(f'부록 A-{b["index"]:02d}. {b["name"]}'),14,True)
            p=doc.add_paragraph();font(p.add_run(b['text']),9)
            table=doc.add_table(rows=2,cols=2);table.autofit=False
            for col in table.columns:col.width=Cm(13.3)
            for i,(label,path) in enumerate(zip(['Stage1 R+E (대표)','Stage2 R+E+F (비교)'],b['files'])):
                font(table.cell(0,i).paragraphs[0].add_run(label),10,True)
                table.cell(1,i).paragraphs[0].add_run().add_picture(str(package/path),width=Cm(13.0))
            p=doc.add_paragraph();font(p.add_run('각 패널: 좌상 Original / 우상 오류 / 좌하 Prediction / 우하 Ground truth. 맞은 픽셀은 흰색, 오류는 정답 색. 100%−image mIoU는 픽셀 오류율이 아니다.'),8)
            md += [f'### A-{b["index"]:02d}. {b["name"]}', '',b['text'],'',*[f'![{label}: {b["name"]}]({path})\n' for label,path in zip(['Stage1','Stage2'],b['files'])],'']
    md+=['## 인용 링크','']+[f'[{k}]: {v["url"]}' for k,v in refs.items()]
    (package/'모델링_결과서_Fusion805.md').write_text('\n'.join(md)+'\n')
    doc.save(package/'모델링_결과서_Fusion805.docx')
    (package/'tables/catalog.json').write_text(json.dumps(table_catalog,ensure_ascii=False,indent=2)+'\n')
    (package/'validation/manuscript_counts.json').write_text(json.dumps(dict(main_figures=fig_counter,tables=table_counter,appendix_pages=20,appendix_images=40,blocks=len(blocks)),indent=2)+'\n')
